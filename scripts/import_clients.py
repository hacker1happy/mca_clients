"""Import shareholder/debenture-holder Excel data into Oracle tables.

The script scans a root folder containing company folders, extracts zip files,
registers files in FILES, validates matching Excel workbooks, and inserts rows
from the "Sheet-With Validations" sheet into CLIENTS.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
import zipfile
from contextlib import closing
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

try:
    import oracledb
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "Missing dependency. Install dependencies with: pip install -r requirements.txt"
    ) from exc


TARGET_NAME_FRAGMENT = "list of share holders debenture holders"
TARGET_EXTENSION = ".xlsm"
TARGET_SHEET = "Sheet-With Validations"

REQUIRED_COLUMNS = {
    "NAME_OF_SHAREHOLDER": "Name of shareholder/ debenture holder",
    "FOLIO_REFERENCE_NUMBER": "Folio Number / Reference Number",
    "DP_CLIENT_ACCOUNT": "DP ID-Client Id-Account",
    "TYPE_OF_IDENTIFIER": "Type of Identifier",
    "IDENTIFICATION_NO": "Identification No.",
    "OCCUPATION": "Occupation",
    "NUMBER_OF_SECURITY_HELD": "Number of security held",
    "TOTAL_SECURITY_AMOUNT_INR": "Total amount of securities held (in INR)",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Import MCA shareholder/debenture-holder Excel data."
    )
    parser.add_argument("root_folder", help="Root folder containing company folders")
    parser.add_argument("--db-user", required=True, help="Oracle database username")
    parser.add_argument("--db-password", required=True, help="Oracle database password")
    parser.add_argument(
        "--db-dsn",
        required=True,
        help="Oracle DSN, for example localhost:1521/FREEPDB1",
    )
    parser.add_argument(
        "--company-name-column",
        default="COMPANY_NAME",
        help="Column in COMPANIES whose value matches the company folder name",
    )
    parser.add_argument(
        "--error-report",
        default="skipped_files_report.json",
        help="Path for the final JSON skipped-file report",
    )
    parser.add_argument(
        "--log-file",
        default="import_clients.log",
        help="Path for importer log output",
    )
    return parser.parse_args()


def configure_logging(log_file: str) -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.FileHandler(log_file), logging.StreamHandler(sys.stdout)],
    )


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def validate_identifier(identifier: str, label: str) -> str:
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", identifier):
        raise ValueError(f"Invalid {label}: {identifier}")
    return identifier.upper()


def extract_zip(zip_path: Path, skipped: list[dict[str, str]]) -> None:
    try:
        with zipfile.ZipFile(zip_path) as archive:
            archive.extractall(zip_path.parent)
        zip_path.unlink()
        logging.info("Extracted and deleted zip: %s", zip_path)
    except Exception as exc:
        message = f"Unable to extract zip file: {exc}"
        skipped.append({"file_path": str(zip_path), "reason": message})
        logging.exception(message)


def extract_all_zips(root_folder: Path, skipped: list[dict[str, str]]) -> None:
    while True:
        zip_files = [path for path in root_folder.rglob("*.zip") if path.is_file()]
        if not zip_files:
            return
        for zip_path in zip_files:
            extract_zip(zip_path, skipped)


def get_company_folder(root_folder: Path, file_path: Path) -> str | None:
    try:
        relative_path = file_path.relative_to(root_folder)
    except ValueError:
        return None
    return relative_path.parts[0] if len(relative_path.parts) > 1 else None


def get_company_number(
    cursor: Any, company_folder: str, company_name_column: str
) -> int | None:
    sql = (
        f"SELECT SEQ_NO FROM COMPANIES "
        f"WHERE UPPER(TRIM({company_name_column})) = UPPER(TRIM(:company_name))"
    )
    cursor.execute(sql, company_name=company_folder)
    row = cursor.fetchone()
    return int(row[0]) if row else None


def get_existing_file_number(cursor: Any, file_name: str, company_number: int) -> int | None:
    cursor.execute(
        """
        SELECT FILE_NUMBER
        FROM FILES
        WHERE COMPANY_NUMBER = :company_number
          AND FILE_NAME = :file_name
        """,
        company_number=company_number,
        file_name=file_name,
    )
    row = cursor.fetchone()
    return int(row[0]) if row else None


def create_file_entry(cursor: Any, file_name: str, company_number: int) -> int:
    file_number_var = cursor.var(oracledb.NUMBER)
    cursor.execute(
        """
        INSERT INTO FILES (FILE_NAME, COMPANY_NUMBER)
        VALUES (:file_name, :company_number)
        RETURNING FILE_NUMBER INTO :file_number
        """,
        file_name=file_name,
        company_number=company_number,
        file_number=file_number_var,
    )
    return int(file_number_var.getvalue()[0])


def get_or_create_file(
    connection: Any, cursor: Any, file_path: Path, company_number: int
) -> tuple[int, bool]:
    existing_file_number = get_existing_file_number(
        cursor, file_path.name, company_number
    )
    if existing_file_number is not None:
        return existing_file_number, False

    file_number = create_file_entry(cursor, file_path.name, company_number)
    connection.commit()
    logging.info("Created FILES entry %s for %s", file_number, file_path)
    return file_number, True


def validate_candidate_workbook(file_path: Path) -> str | None:
    if TARGET_NAME_FRAGMENT not in file_path.name.lower():
        return "File name does not contain 'List of share holders debenture holders'"
    if file_path.suffix.lower() != TARGET_EXTENSION:
        return "File extension is not .xlsm"
    return None


def locate_headers(sheet: Any) -> dict[str, int]:
    expected = {
        target_column: normalize_header(source_header)
        for target_column, source_header in REQUIRED_COLUMNS.items()
    }

    for row in sheet.iter_rows(min_row=1, max_row=30):
        available = {
            normalize_header(cell.value): cell.column for cell in row if cell.value
        }
        if all(header in available for header in expected.values()):
            return {
                target_column: available[normalized_header]
                for target_column, normalized_header in expected.items()
            } | {"__HEADER_ROW__": row[0].row}

    missing_headers = ", ".join(REQUIRED_COLUMNS.values())
    raise ValueError(f"Required column headers not found: {missing_headers}")


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def to_decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, Decimal):
        return value
    cleaned = str(value).replace(",", "").strip()
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid numeric value: {value}") from exc


def read_client_rows(file_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, read_only=True, data_only=True, keep_vba=False)
    try:
        if TARGET_SHEET not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {TARGET_SHEET}")

        sheet = workbook[TARGET_SHEET]
        header_map = locate_headers(sheet)
        header_row = header_map.pop("__HEADER_ROW__")
        rows: list[dict[str, Any]] = []

        for row_number in range(header_row + 1, sheet.max_row + 1):
            row_data = {
                target_column: sheet.cell(row=row_number, column=column_number).value
                for target_column, column_number in header_map.items()
            }

            if all(value is None or str(value).strip() == "" for value in row_data.values()):
                continue

            rows.append(
                {
                    "NAME_OF_SHAREHOLDER": clean_text(row_data["NAME_OF_SHAREHOLDER"]),
                    "FOLIO_REFERENCE_NUMBER": clean_text(
                        row_data["FOLIO_REFERENCE_NUMBER"]
                    ),
                    "DP_CLIENT_ACCOUNT": clean_text(row_data["DP_CLIENT_ACCOUNT"]),
                    "TYPE_OF_IDENTIFIER": clean_text(row_data["TYPE_OF_IDENTIFIER"]),
                    "IDENTIFICATION_NO": clean_text(row_data["IDENTIFICATION_NO"]),
                    "OCCUPATION": clean_text(row_data["OCCUPATION"]),
                    "NUMBER_OF_SECURITY_HELD": to_decimal(
                        row_data["NUMBER_OF_SECURITY_HELD"]
                    ),
                    "TOTAL_SECURITY_AMOUNT_INR": to_decimal(
                        row_data["TOTAL_SECURITY_AMOUNT_INR"]
                    ),
                }
            )

        return rows
    finally:
        workbook.close()


def insert_client_rows(cursor: Any, file_number: int, rows: list[dict[str, Any]]) -> None:
    for row in rows:
        row["FILE_NUMBER"] = file_number

    cursor.executemany(
        """
        INSERT INTO CLIENTS (
          NAME_OF_SHAREHOLDER,
          FOLIO_REFERENCE_NUMBER,
          DP_CLIENT_ACCOUNT,
          TYPE_OF_IDENTIFIER,
          IDENTIFICATION_NO,
          OCCUPATION,
          NUMBER_OF_SECURITY_HELD,
          TOTAL_SECURITY_AMOUNT_INR,
          FILE_NUMBER
        )
        VALUES (
          :NAME_OF_SHAREHOLDER,
          :FOLIO_REFERENCE_NUMBER,
          :DP_CLIENT_ACCOUNT,
          :TYPE_OF_IDENTIFIER,
          :IDENTIFICATION_NO,
          :OCCUPATION,
          :NUMBER_OF_SECURITY_HELD,
          :TOTAL_SECURITY_AMOUNT_INR,
          :FILE_NUMBER
        )
        """,
        rows,
    )


def add_skipped(skipped: list[dict[str, str]], file_path: Path, reason: str) -> None:
    skipped.append({"file_path": str(file_path), "reason": reason})
    logging.warning("Skipped %s: %s", file_path, reason)


def process_file(
    connection: Any,
    cursor: Any,
    root_folder: Path,
    file_path: Path,
    company_name_column: str,
    skipped: list[dict[str, str]],
) -> None:
    company_folder = get_company_folder(root_folder, file_path)
    if company_folder is None:
        add_skipped(skipped, file_path, "File is not inside a company folder")
        return

    company_number = get_company_number(cursor, company_folder, company_name_column)
    if company_number is None:
        add_skipped(
            skipped,
            file_path,
            f"Company folder '{company_folder}' was not found in COMPANIES",
        )
        return

    try:
        file_number, created = get_or_create_file(
            connection, cursor, file_path, company_number
        )
    except Exception as exc:
        connection.rollback()
        add_skipped(skipped, file_path, f"Unable to create FILES entry: {exc}")
        return

    if not created:
        logging.info("Duplicate file already exists in FILES, skipping import: %s", file_path)
        return

    validation_error = validate_candidate_workbook(file_path)
    if validation_error:
        add_skipped(skipped, file_path, validation_error)
        return

    try:
        client_rows = read_client_rows(file_path)
        if not client_rows:
            raise ValueError("No client rows found in the validation sheet")
        insert_client_rows(cursor, file_number, client_rows)
        connection.commit()
        logging.info(
            "Imported %s client rows from %s into FILE_NUMBER %s",
            len(client_rows),
            file_path,
            file_number,
        )
    except Exception as exc:
        connection.rollback()
        add_skipped(skipped, file_path, f"Unable to import workbook: {exc}")


def write_error_report(report_path: Path, skipped: list[dict[str, str]]) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps({"skipped_files": skipped}, indent=2),
        encoding="utf-8",
    )
    logging.info("Wrote skipped-file report: %s", report_path)


def main() -> int:
    args = parse_args()
    configure_logging(args.log_file)

    root_folder = Path(args.root_folder).expanduser().resolve()
    if not root_folder.is_dir():
        logging.error("Root folder does not exist or is not a directory: %s", root_folder)
        return 1

    try:
        company_name_column = validate_identifier(
            args.company_name_column, "company name column"
        )
    except ValueError as exc:
        logging.error(str(exc))
        return 1

    skipped: list[dict[str, str]] = []
    extract_all_zips(root_folder, skipped)

    files = sorted(path for path in root_folder.rglob("*") if path.is_file())
    logging.info("Scanning %s files under %s", len(files), root_folder)

    with closing(
        oracledb.connect(
            user=args.db_user,
            password=args.db_password,
            dsn=args.db_dsn,
        )
    ) as connection:
        with closing(connection.cursor()) as cursor:
            for file_path in files:
                process_file(
                    connection,
                    cursor,
                    root_folder,
                    file_path,
                    company_name_column,
                    skipped,
                )

    write_error_report(Path(args.error_report), skipped)
    logging.info("Import completed with %s skipped files", len(skipped))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
