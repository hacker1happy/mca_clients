import os
import shutil
import logging
from datetime import datetime
from decimal import Decimal, InvalidOperation

import openpyxl
import oracledb
from tqdm.notebook import tqdm


TARGET_TEXT = "list of share holders debenture holders"
REQUIRED_SHEETS = ("Sheet-With Validations", "Sheet-Without Validations")
COLUMN_MAP = {
    "NAME_OF_SHAREHOLDER": 0,          # E
    "FOLIO_NUMBER": 3,                 # H
    "DP_CLIENT_ACCOUNT": 4,            # I
    "TYPE_OF_IDENTIFIER": 7,           # L
    "IDENTIFICATION_NO": 8,            # M
    "OCCUPATION": 9,                   # N
    "NUMBER_OF_SECURITY_HELD": 10,     # O
    "TOTAL_SECURITY_AMOUNT_INR": 12,   # Q
}

DB_CONFIG = {
    "user": "KINHERITANCE_USER",
    "password": "OurStartup1!",
    "dsn": """
    (description=
      (retry_count=20)
      (retry_delay=3)
      (address=
        (protocol=tcps)
        (port=1522)
        (host=adb.ap-mumbai-1.oraclecloud.com)
      )
      (connect_data=
        (service_name=gbad14cef94c10b_shareholdersdetails_medium.adb.oraclecloud.com)
      )
      (security=
        (ssl_server_dn_match=yes)
      )
    )
    """
}


def setup_logging(root_folder):
    log_dir = os.path.join(root_folder, "logs")
    os.makedirs(log_dir, exist_ok=True)
    log_file = os.path.join(log_dir, f"mca_file_migration_{datetime.now():%Y%m%d_%H%M%S}.log")

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[logging.FileHandler(log_file)],
        force=True,
    )
    logging.info("Log file: %s", log_file)
    return log_file


def unique_destination(path):
    if not os.path.exists(path):
        return path
    suffix = datetime.now().strftime("%Y%m%d_%H%M%S")
    parent_dir = os.path.dirname(path)
    name = os.path.basename(path)
    if os.path.isdir(path):
        return os.path.join(parent_dir, f"{name}_{suffix}")
    stem, ext = os.path.splitext(name)
    return os.path.join(parent_dir, f"{stem}_{suffix}{ext}")


def move_path(source, target_root, relative_path):
    destination = unique_destination(os.path.join(target_root, relative_path))
    os.makedirs(os.path.dirname(destination), exist_ok=True)
    shutil.move(source, destination)
    return destination


def find_company(cursor, folder_name):
    cursor.execute(
        """
        SELECT COMP_ID
        FROM COMPANIES
        WHERE UPPER(TRIM(FOLDER_NAME)) = UPPER(TRIM(:folder_name))
        """,
        folder_name=folder_name,
    )
    row = cursor.fetchone()
    return row[0] if row else None


def get_or_create_file(cursor, company_id, relative_file_name):
    # FILE_NAME stores the path below the company folder to avoid false duplicates.
    cursor.execute(
        """
        SELECT FILE_ID
        FROM FILES
        WHERE COMPANY_ID = :company_id
          AND FILE_NAME = :file_name
        """,
        company_id=company_id,
        file_name=relative_file_name,
    )
    row = cursor.fetchone()
    if row:
        return row[0]

    file_id_var = cursor.var(oracledb.NUMBER)
    try:
        cursor.execute(
            """
            INSERT INTO FILES (FILE_NAME, COMPANY_ID, CREATED_AT)
            VALUES (:file_name, :company_id, SYSTIMESTAMP)
            RETURNING FILE_ID INTO :file_id
            """,
            file_name=relative_file_name,
            company_id=company_id,
            file_id=file_id_var,
        )
    except oracledb.IntegrityError:
        cursor.execute(
            """
            SELECT FILE_ID
            FROM FILES
            WHERE COMPANY_ID = :company_id
              AND FILE_NAME = :file_name
            """,
            company_id=company_id,
            file_name=relative_file_name,
        )
        return cursor.fetchone()[0]
    return int(file_id_var.getvalue()[0])


def validate_file(file_path):
    file_name = os.path.basename(file_path)
    if TARGET_TEXT not in file_name.lower():
        raise ValueError("Filename does not contain required text")
    if os.path.splitext(file_name)[1].lower() != ".xlsm":
        raise ValueError("File extension is not .xlsm")

    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_vba=True)
    missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in workbook.sheetnames]
    if missing:
        workbook.close()
        raise ValueError(f"Missing required sheet(s): {', '.join(missing)}")
    return workbook


def clean_text(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def clean_number(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError):
        return None


def sheet_rows(workbook):
    # Headers are assumed to be row 1; MCA data starts at row 2.
    rows = []
    for sheet_name in REQUIRED_SHEETS:
        sheet = workbook[sheet_name]
        for source_row in sheet.iter_rows(min_row=2, min_col=5, max_col=17, values_only=True):
            values = {key: source_row[index] for key, index in COLUMN_MAP.items()}
            if not any(values.values()):
                continue
            rows.append(
                (
                    clean_text(values["NAME_OF_SHAREHOLDER"]),
                    clean_text(values["FOLIO_NUMBER"]),
                    clean_text(values["DP_CLIENT_ACCOUNT"]),
                    clean_text(values["TYPE_OF_IDENTIFIER"]),
                    clean_text(values["IDENTIFICATION_NO"]),
                    clean_text(values["OCCUPATION"]),
                    clean_number(values["NUMBER_OF_SECURITY_HELD"]),
                    clean_number(values["TOTAL_SECURITY_AMOUNT_INR"]),
                )
            )
    return rows


def reload_clients(cursor, file_id, client_rows):
    # Reprocessing is resume-safe: old rows for this file are replaced.
    cursor.execute("DELETE FROM MCA_CLIENTS WHERE FILE_ID = :file_id", file_id=file_id)
    if not client_rows:
        return

    cursor.executemany(
        """
        INSERT INTO MCA_CLIENTS (
            NAME_OF_SHAREHOLDER,
            FOLIO_NUMBER,
            DP_CLIENT_ACCOUNT,
            TYPE_OF_IDENTIFIER,
            IDENTIFICATION_NO,
            OCCUPATION,
            NUMBER_OF_SECURITY_HELD,
            TOTAL_SECURITY_AMOUNT_INR,
            FILE_ID
        )
        VALUES (:1, :2, :3, :4, :5, :6, :7, :8, :9)
        """,
        [row + (file_id,) for row in client_rows],
    )


def process_file(conn, cursor, company_id, company_folder, file_path, root_folder):
    relative_path = os.path.relpath(file_path, company_folder)
    relative_name = relative_path.replace(os.sep, "/")
    workbook = None
    try:
        workbook = validate_file(file_path)
        rows = sheet_rows(workbook)
        file_id = get_or_create_file(cursor, company_id, relative_name)
        reload_clients(cursor, file_id, rows)
        conn.commit()
        logging.info("Processed %s | FILE_ID=%s | rows=%s", file_path, file_id, len(rows))
    except Exception as exc:
        conn.rollback()
        company_name = os.path.basename(company_folder)
        destination = move_path(
            file_path,
            os.path.join(root_folder, "unprocessed"),
            os.path.join(company_name, relative_path),
        )
        logging.exception("Skipped %s | moved to %s | error=%s", file_path, destination, exc)
    finally:
        if workbook:
            workbook.close()


def process_company(conn, cursor, company_folder, root_folder):
    company_name = os.path.basename(company_folder)
    company_id = find_company(cursor, company_name)
    if not company_id:
        destination = move_path(company_folder, os.path.join(root_folder, "company_not_found"), company_name)
        logging.warning("Company not found: %s | moved to %s", company_name, destination)
        return

    # Collect first so moving invalid files does not disturb directory walking.
    files = []
    for folder_path, _, file_names in os.walk(company_folder):
        for file_name in file_names:
            files.append(os.path.join(folder_path, file_name))
    for file_path in tqdm(files, desc=f"Files: {company_name}", unit="file"):
        process_file(conn, cursor, company_id, company_folder, file_path, root_folder)


def main():
    root_folder = r'keep root folder path here'
    source_root = os.path.join(root_folder, "Files_to_be_moved_in_DB")
    log_file = setup_logging(root_folder)
    logging.info("Migration started | root_folder=%s", root_folder)

    if not os.path.isdir(source_root):
        raise NotADirectoryError(f"Missing folder: {source_root}")

    conn = None
    cursor = None
    try:
        conn = oracledb.connect(**DB_CONFIG)
        cursor = conn.cursor()
        company_folders = [
            os.path.join(source_root, name)
            for name in os.listdir(source_root)
            if os.path.isdir(os.path.join(source_root, name))
        ]
        for company_folder in tqdm(sorted(company_folders), desc="Companies", unit="company"):
            process_company(conn, cursor, company_folder, root_folder)
        logging.info("Migration completed successfully | log_file=%s", log_file)
    except Exception:
        if conn:
            conn.rollback()
        logging.exception("Migration failed")
        raise
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


if __name__ == "__main__":
    main()
